"""Build course slides, speaker notes and a searchable handbook from the video source."""

import argparse
import hashlib
import html
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from build_course import ROOT, read
from build_course_video import DESTINATION, timecode
from narrate_course_video import SOURCE, digest

SLIDES = 'NetMambaPlus-course-slides.pptx'
SLIDE_PDF = 'NetMambaPlus-course-slides.pdf'
HANDBOOK = 'NetMambaPlus-course-handbook.pdf'
REPO = 'https://github.com/buffbeefalo/netmambaplus-reproduction/blob/main/'


def paragraph_text(text, base=ROOT / 'docs'):
    """Preserve useful Markdown links when printing the file inventory."""
    links = []
    def save(match):
        label, target = match.groups()
        if target.startswith('#'):
            target = REPO + 'docs/repository-walkthrough.md' + target
        elif not target.startswith(('https://', 'http://')):
            relative, separator, fragment = target.partition('#')
            path = (base / relative).resolve()
            if not path.is_relative_to(ROOT):
                raise ValueError('Document link escapes the repository')
            target = REPO + path.relative_to(ROOT).as_posix() + (separator + fragment if separator else '')
        links.append(f'<link href="{html.escape(target, quote=True)}" color="#126e68">{html.escape(label)}</link>')
        return f'LINKPLACEHOLDER{len(links) - 1}END'
    text = re.sub(r'<a id="[^"]+"></a>', '', text)
    rendered = html.escape(re.sub(r'\[([^]]+)\]\(([^)]+)\)', save, text))
    rendered = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', rendered)
    rendered = rendered.replace('`', '')
    for i, link in enumerate(links):
        rendered = rendered.replace(f'LINKPLACEHOLDER{i}END', link)
    return rendered


def build(work, output):
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, PageBreak, Table, TableStyle, Preformatted

    manifest = read(output / 'media-manifest.json')
    source_hash = digest(SOURCE)
    guide_hash = digest(ROOT / 'docs/repository-walkthrough.md')
    if manifest['source_sha256'] != source_hash:
        raise ValueError('Render the current lesson video before building its documents')
    frames = []
    for scene in manifest['scenes']:
        path = work / 'frames' / (scene['id'] + '.png')
        if not path.is_file():
            raise ValueError('Missing rendered lesson frame: ' + scene['id'])
        expected = scene.get('visuals', [])
        if len(expected) != 1 or digest(path) != expected[0]['sha256']:
            raise ValueError('Slide frame differs from the video build: ' + scene['id'])
        if digest(ROOT / expected[0]['path']) != expected[0]['sha256']:
            raise ValueError('Published slide frame differs from the video build: ' + scene['id'])
        frames.append(path)
    output.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = 'NetMamba+ — complete repository course'
    deck.core_properties.author = 'Codex'
    deck.core_properties.last_modified_by = 'Codex'
    built_at = datetime.now(timezone.utc)
    deck.core_properties.created = deck.core_properties.modified = built_at
    pdf_date = built_at.strftime("D:%Y%m%d%H%M%S") + "+00'00'"
    deck.core_properties.subject = 'Source-bound course slides with editable narration notes'
    slide_pdf = canvas.Canvas(str(output / SLIDE_PDF), pagesize=(960, 540))
    slide_pdf.setTitle(deck.core_properties.title)
    slide_pdf.setAuthor('Codex')
    slide_pdf.setDateFormatter(lambda *parts: pdf_date)
    for scene, frame in zip(manifest['scenes'], frames):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        picture = slide.shapes.add_picture(str(frame), 0, 0, width=deck.slide_width, height=deck.slide_height)
        picture._element.nvPicPr.cNvPr.set('descr', scene['title'] + ': ' + ' '.join(scene['bullets']))
        body = scene['narration'] or f"Practice for {scene['seconds']} seconds: " + ' '.join(scene['bullets'])
        references = '\n'.join(REPO + manifest['references'][key]['path'] for key in scene['references'])
        slide.notes_slide.notes_text_frame.text = (f"{timecode(scene['start'])[:8]} — {scene['title']}\n\n{body}\n\nEvidence:\n{references}\n\nSlide image follows the video; notes are editable. Source SHA-256: {manifest['source_sha256']}")
        slide_pdf.bookmarkPage(scene['id'])
        slide_pdf.addOutlineEntry(scene['title'].replace('\n', ': '), scene['id'], level=0)
        slide_pdf.drawImage(str(frame), 0, 0, width=960, height=540)
        slide_pdf.showPage()
    deck.save(output / SLIDES)
    slide_pdf.save()

    fonts = Path('/usr/share/fonts/truetype/dejavu')
    for name, file in [('Course', 'DejaVuSans.ttf'), ('CourseBold', 'DejaVuSans-Bold.ttf'), ('CourseMono', 'DejaVuSansMono.ttf')]:
        pdfmetrics.registerFont(TTFont(name, str(fonts / file)))
    pdfmetrics.registerFontFamily('Course', normal='Course', bold='CourseBold', italic='Course', boldItalic='CourseBold')
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle('CourseBody', fontName='Course', fontSize=10, leading=15, spaceAfter=9, textColor=colors.HexColor('#172c3a'))
    small_style = ParagraphStyle('CourseSmall', parent=body_style, fontSize=8, leading=11)
    heading_style = ParagraphStyle('CourseHeading', parent=body_style, fontName='CourseBold', fontSize=19, leading=25, spaceBefore=12, keepWithNext=True)
    subheading_style = ParagraphStyle('CourseSubheading', parent=body_style, fontName='CourseBold', fontSize=12, leading=17, spaceBefore=10, keepWithNext=True)
    code_style = ParagraphStyle('CourseCode', parent=body_style, fontName='CourseMono', fontSize=7, leading=10)
    story = [Paragraph('NetMamba+<br/>The complete repository course', heading_style),
             Paragraph('Codex-authored handbook, narration script and every-file appendix', body_style),
             Paragraph('This handbook follows the single hour-long video and its slide deck. The video uses synthetic narration. Automated checks and their limits are recorded separately; no complete human watch-through is claimed.', body_style),
             Paragraph('Measured three-seed mean: <b>86.65%</b>. Paper result: <b>97.50%, not reproduced</b>. The demo replays saved predictions. Live IDS and NPU/SmartNIC deployment remain future work.', body_style),
             Paragraph('Slide images in the PowerPoint match the video; their narration is available as editable speaker notes. This handbook provides searchable text and a complete repository-file appendix. The original customer presentation remains a separate historical package.', body_style),
             Paragraph('Source SHA-256: ' + manifest['source_sha256'], small_style),
             Paragraph('Course chapters', subheading_style)]
    for c in manifest['chapters']:
        story.append(Paragraph(timecode(c['start'])[:8] + ' — ' + html.escape(c['title']), body_style))
    for number, chapter in enumerate(manifest['chapters'], 1):
        story += [PageBreak(), Paragraph(f"{number}. {html.escape(chapter['title'])}", heading_style)]
        for scene in [s for s in manifest['scenes'] if s['chapter'] == number]:
            story.append(Paragraph(timecode(scene['start'])[:8] + ' — ' + html.escape(scene['title']), subheading_style))
            if scene['narration']:
                story.append(Paragraph(html.escape(scene['narration']), body_style))
            else:
                story.append(Paragraph(f"Practice for {scene['seconds']} seconds: " + html.escape(' '.join(scene['bullets'])), body_style))
            story.append(Paragraph('<b>On screen:</b> ' + html.escape(' · '.join(scene['bullets'])), small_style))
            if 'code' in scene:
                story.append(Preformatted(scene['code'], code_style, maxLineLength=90))
            if 'rows' in scene:
                for row in [scene['columns']] + scene['rows']:
                    story.append(Paragraph(' — '.join(html.escape(str(c)) for c in row), small_style))
            refs = ', '.join(f'<link href="{REPO + manifest["references"][key]["path"]}" color="#126e68">{html.escape(key)}</link>' for key in scene['references'])
            story.append(Paragraph('Evidence: ' + refs, small_style))
    guide = ROOT / 'docs/repository-walkthrough.md'
    story += [PageBreak(), Paragraph('Appendix: every file and every download', heading_style),
              Paragraph('The following reviewed file guide is included so the entire handoff remains readable offline. File links point to the repository; downloaded model inputs and weights are described separately.', body_style)]
    table_rows, code_lines, in_code = [], [], False
    def flush_table():
        if not table_rows:
            return
        columns = len(table_rows[0])
        widths = [155, 360] if columns == 2 else [515 / columns] * columns
        data = [[Paragraph(paragraph_text(c), small_style) for c in row] for row in table_rows]
        t = Table(data, colWidths=widths, repeatRows=1, hAlign='LEFT')
        t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e2eeeb')),('LINEBELOW',(0,0),(-1,0),0.5,colors.HexColor('#126e68')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        story.extend([t, Spacer(1, 10)])
        table_rows.clear()
    for line in guide.read_text(encoding='utf-8').splitlines():
        if line.startswith('```'):
            flush_table()
            if in_code:
                story.append(Preformatted('\n'.join(code_lines), code_style, maxLineLength=95))
                code_lines.clear()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith('|'):
            if not re.fullmatch(r'[|: \-]+', line):
                table_rows.append([c.strip() for c in line.strip('|').split('|')])
            continue
        flush_table()
        if not line.strip():
            continue
        level = len(line) - len(line.lstrip('#'))
        style = heading_style if level in (1, 2) else subheading_style if level else body_style
        text = line[level:].strip() if level else line
        story.append(Paragraph(paragraph_text(text), style))
    flush_table()
    def footer(canv, doc):
        canv.setDateFormatter(lambda *parts: pdf_date)
        canv.setFont('Course', 8)
        canv.drawString(40, 22, 'NetMamba+ · Codex course · measured reproduction attempt')
        canv.drawRightString(A4[0] - 40, 22, str(doc.page))
    doc = SimpleDocTemplate(str(output / HANDBOOK), pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40,
                            title='NetMamba+ course handbook and complete file guide', author='Codex')
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    if digest(SOURCE) != source_hash or digest(guide) != guide_hash:
        raise ValueError('Lesson or walkthrough changed while documents were generated; rebuild before publishing')
    record = {'schema_version':1, 'source_sha256':source_hash, 'walkthrough_sha256':guide_hash,
              'slide_count':len(frames), 'spoken_scenes':sum(bool(s['narration']) for s in manifest['scenes']),
              'slide_images':{p.name:digest(p) for p in frames},
              'frames':[{'scene_id':s['id'], 'path':(output/'frames'/p.name).relative_to(ROOT).as_posix(), 'sha256':digest(p)} for s,p in zip(manifest['scenes'],frames)],
              'artifacts':{name:{'bytes':(output/name).stat().st_size,'sha256':digest(output/name)} for name in [SLIDES, SLIDE_PDF, HANDBOOK]},
              'scope':'Generated slides share the rendered frames; speaker notes and searchable handbook share narration. Handbook appends the reviewed complete file guide. Separate artifact checks are required.'}
    (output / 'companion-manifest.json').write_text(json.dumps(record, indent=2) + '\n', encoding='utf-8', newline='\n')
    print(json.dumps({'slides':len(frames), 'artifacts':record['artifacts']}, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--work-dir', type=Path, default=ROOT/'runs/video-course/v3-neural')
    parser.add_argument('--output', type=Path, default=DESTINATION)
    args = parser.parse_args()
    build(args.work_dir.resolve(), args.output.resolve())


if __name__ == '__main__':
    main()
