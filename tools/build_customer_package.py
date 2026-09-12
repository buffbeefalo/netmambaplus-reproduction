"""Build the editable deck, briefing PDF and charts from reviewed measurements."""

import html
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
CUSTOMER = ROOT / "docs/customer"
EVIDENCE = CUSTOMER / "evidence"
REPO = "https://github.com/buffbeefalo/netmambaplus-reproduction"
INK, PAPER, TEAL, AMBER, MUTED = "132C38", "F4F1E8", "007C78", "A34F17", "516774"


def load(name):
    return json.loads((EVIDENCE / name).read_text())


def percent(value):
    return f"{100 * value:.2f}%"


def data_and_tokens():
    results = load("results.json")
    benchmark = load("benchmark/metrics.json")
    export = load("export-probe/metrics.json")
    if [row["seed"] for row in results["seeds"]] != [0, 1, 2]:
        raise ValueError("The customer package requires all three declared seed results")
    if any(row["completed_epochs"] != 120 or not row["fresh_strict_metrics_agree"] for row in results["seeds"]):
        raise ValueError("Training or strict inference evidence is incomplete")
    rows = [[str(row["seed"]), percent(row["metrics"]["accuracy"]), percent(row["metrics"]["weighted_f1"]),
             percent(row["metrics"]["macro_f1"])] for row in results["seeds"]]
    aggregate = results["aggregate"]
    rows.append(["Mean", *(percent(aggregate[key]["mean"]) for key in ("accuracy", "weighted_f1", "macro_f1"))])
    table = "| Seed | Accuracy | Weighted F1 | Macro F1 |\n|---|---|---|---|\n"
    table += "\n".join("| " + " | ".join(row) + " |" for row in rows)
    latencies = [[str(row["batch_size"]), f"{row['p50_ms']:.2f} ms", f"{row['p95_ms']:.2f} ms",
                  f"{row['flows_per_second_from_mean']:,.1f}"] for row in benchmark["measurements"]]
    latency_table = "| Batch | Median latency | 95th percentile | Flows/s from mean |\n|---|---|---|---|\n"
    latency_table += "\n".join("| " + " | ".join(row) + " |" for row in latencies)
    tokens = {"mean_accuracy": percent(aggregate["accuracy"]["mean"]),
              "mean_weighted_f1": percent(aggregate["weighted_f1"]["mean"]),
              "mean_macro_f1": percent(aggregate["macro_f1"]["mean"]),
              "results_table": table, "latency_table": latency_table,
              "export_status": "captured" if export["export_status"] == "captured" else "not captured in the tested Torch path"}
    return results, benchmark, tokens, rows, latencies


def substitute(text, tokens):
    for name, value in tokens.items():
        text = text.replace("@@" + name + "@@", value)
    if "@@" in text:
        raise ValueError("Unresolved presentation token")
    return text


def charts(results):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    output = CUSTOMER / "figures"
    output.mkdir(exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "figure.facecolor": "#" + PAPER, "axes.facecolor": "#" + PAPER})
    figure, axes = plt.subplots(1, 2, figsize=(11.8, 4.2))
    for row, color in zip(results["seeds"], ("#007C78", "#325C96", "#A34F17")):
        history = row["curve"]
        epochs = [item["epoch"] + 1 for item in history]
        axes[0].plot(epochs, [item["train_loss"] for item in history], color=color, label=f"Seed {row['seed']}", lw=1.6)
        axes[1].plot(epochs, [100 * item["valid_acc"] for item in history], color=color, label=f"Seed {row['seed']}", lw=1.6)
    axes[0].set(ylabel="Training loss", xlabel="Completed epoch", title="Supervised training history")
    axes[1].set(ylabel="Validation accuracy (%)", xlabel="Completed epoch", title="Checkpoint selection uses validation")
    for axis in axes:
        axis.grid(alpha=.2)
        axis.legend(frameon=False, fontsize=9)
    figure.tight_layout()
    for extension in ("png", "svg", "pdf"):
        figure.savefig(output / f"learning-curves.{extension}", dpi=160, bbox_inches="tight")
    plt.close(figure)
    matrix = np.array(results["seeds"][0]["metrics"]["confusion_matrix"])
    names = ["Flood", "RTSP brute", "Audio", "Other", "Cameras", "Home auto"]
    figure, axis = plt.subplots(figsize=(7.7, 4.9))
    display = axis.imshow(matrix, cmap="GnBu", vmin=0, vmax=max(1, int(matrix.max())))
    axis.set(xticks=range(6), yticks=range(6), xticklabels=names, yticklabels=names,
             xlabel="Predicted class", ylabel="True class")
    for r in range(6):
        for c in range(6):
            axis.text(c, r, str(matrix[r, c]), ha="center", va="center", fontsize=12,
                      color="white" if matrix[r, c] > matrix.max() * .55 else "#" + INK)
    figure.colorbar(display, ax=axis, label="Test flows", fraction=.046, pad=.04)
    figure.tight_layout()
    for extension in ("png", "svg", "pdf"):
        figure.savefig(output / f"seed0-confusion.{extension}", dpi=180, bbox_inches="tight")
    plt.close(figure)


def build_slides(results, benchmark, tokens, score_rows, latency_rows):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    source = json.loads(substitute((CUSTOMER / "presentation-source.json").read_text(), tokens))
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    deck.core_properties.title = source["title"]
    deck.core_properties.author = "NetMamba+ reproduction project"
    deck.core_properties.subject = "Measured research, customer explanation and deployment boundaries"
    notes = ["# Presenter talk track", "", "Use slides 1–8 for the core walkthrough; keep the remaining technical detail available for questions.", ""]

    def rectangle(slide, x, y, width, height, fill=PAPER, line=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line:
            shape.line.color.rgb = RGBColor.from_string(line)
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def text(slide, content, x, y, width, height, size=18, color=INK, bold=False, font="DejaVu Sans"):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.02)
        frame.margin_top = frame.margin_bottom = 0
        frame.text = content
        for paragraph in frame.paragraphs:
            paragraph.font.name, paragraph.font.size = font, Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(9)
        return shape

    def cards(slide, values, y=2.65, height=3.22, size=17):
        gap, total, start = .22, 12.03, .65
        width = (total - gap * (len(values) - 1)) / len(values)
        for index, (title, body) in enumerate(values):
            x = start + index * (width + gap)
            rectangle(slide, x, y, width, height, "FFFFFF")
            rectangle(slide, x, y, width, .055, TEAL)
            text(slide, title, x + .22, y + .2, width - .44, .5, 21, TEAL, True)
            text(slide, body, x + .22, y + .92, width - .44, height - 1.05, size)

    def table(slide, headings, values, y=2.55, widths=None, size=18):
        rows = [headings, *values]
        widths = widths or [12.0 / len(headings)] * len(headings)
        height = .53
        for r, row in enumerate(rows):
            x = .65
            for width, value in zip(widths, row):
                fill = INK if r == 0 else ("FFFFFF" if r % 2 else "E5EBE6")
                rectangle(slide, x, y + r * height, width, height, fill)
                text(slide, value, x + .14, y + r * height + .12, width - .24, .37,
                     size if r else 14, "FFFFFF" if r == 0 else INK, r == 0)
                x += width

    for index, item in enumerate(source["slides"], 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        layout = item["layout"]
        rectangle(slide, 0, 0, 13.333333, 7.5, INK if layout == "cover" else PAPER)
        slide.notes_slide.notes_text_frame.text = item["notes"] + "\n\nEvidence: " + REPO + "/tree/main/docs/customer/evidence"
        notes += [f"## Slide {index} · {item['title'].replace(chr(10), ' ')}", "", item["notes"], ""]
        if layout == "cover":
            text(slide, item["section"].upper(), .7, .55, 11.8, .35, 13, "85DBCB", True)
            text(slide, item["title"], .7, 1.28, 11.8, 1.9, 43, "FFFFFF", font="Liberation Serif")
            text(slide, item["subtitle"], .7, 3.4, 11.65, .65, 19, "CCDBDE")
            for n, (value, label) in enumerate(item["cards"]):
                x = .7 + n * 4.05
                rectangle(slide, x, 4.55, 3.8, 1.35, "203C48")
                text(slide, value, x + .2, 4.68, 3.42, .6, 33, "85DBCB", True)
                text(slide, label, x + .2, 5.39, 3.42, .45, 13, "FFFFFF")
            text(slide, item["takeaway"], .7, 6.4, 11.9, .6, 14, "CCDBDE")
            text(slide, source["date"], .7, 7.05, 5, .2, 10, "CCDBDE")
            continue

        text(slide, item["section"].upper(), .65, .37, 11.8, .3, 11, TEAL, True)
        title_size = 34 if len(item["title"]) < 54 else 31
        text(slide, item["title"], .65, .92, 12.0, .68, title_size, font="Liberation Serif")
        text(slide, item["subtitle"], .65, 1.75, 11.9, .6, 16, MUTED)
        if layout in ("cards", "closing", "evidence", "inputs"):
            cards(slide, item["cards"], height=3.15, size=16 if layout == "evidence" else 17)
        elif layout == "protocol":
            table(slide, ["Setting", "Measured source profile", "Paper v1"], [
                ["Fine-tuning budget", "120 epochs × 3 seeds", "120 epochs"],
                ["Batch size", "128", "64"],
                ["Maximum effective LR", "0.001 (blr 0.002)", "Reported 0.002"],
                ["Runtime", "GB10 / Torch 2.9.1 / cu130", "A100 / Torch 2.1.1"],
                ["Pretraining history", "Released weights; partly unknown", "Browser + Kitsune"]
            ], widths=[3.0, 4.85, 4.15], size=16)
        elif layout == "results":
            table(slide, ["Seed", "Accuracy", "Weighted F1", "Macro F1"], score_rows)
            sd = results["aggregate"]["accuracy"]["sample_standard_deviation"] * 100
            text(slide, f"Accuracy sample SD: {sd:.2f} percentage points across seeds on this fixed split.",
                 .8, 5.48, 11.7, .55, 15, MUTED)
        elif layout == "confusion":
            metrics = results["seeds"][0]["metrics"]
            names = ["Flood", "RTSP", "Audio", "Other", "Camera", "Home"]
            text(slide, "TRUE / PRED.", .67, 2.57, 1.18, .3, 9, MUTED, True)
            for c, name in enumerate(names):
                text(slide, name, 1.95 + c * 1.0, 2.55, .95, .35, 12, MUTED, True)
            for r, name in enumerate(names):
                text(slide, name, .67, 3.05 + r * .5, 1.18, .35, 13, INK, True)
                for c, value in enumerate(metrics["confusion_matrix"][r]):
                    fraction = value / 200
                    start, end = (232, 241, 237), (0, 105, 110)
                    shade = ''.join(f"{round(a + (b - a) * fraction):02X}" for a, b in zip(start, end))
                    rectangle(slide, 1.9 + c * 1.0, 2.94 + r * .5, .94, .46, shade)
                    text(slide, str(value), 2.11 + c * 1.0, 3.02 + r * .5, .64, .35, 19,
                         "FFFFFF" if fraction > .5 else INK)
            cards_data = [("Seed 0 accuracy", percent(metrics["accuracy"])),
                          ("Seed 0 macro F1", percent(metrics["macro_f1"])), ("Test examples", "1,041 flows / six classes")]
            for n, (label, value) in enumerate(cards_data):
                text(slide, label, 8.4, 2.65 + n * 1.04, 3.95, .35, 14, MUTED)
                text(slide, value, 8.4, 3.02 + n * 1.04, 3.95, .6, 26, TEAL, True)
        elif layout == "demo":
            from PIL import Image
            screenshot = CUSTOMER / "demo/replay-screenshot.png"
            if not screenshot.is_file():
                raise ValueError("Actual demo screenshot must exist before building the final deck")
            with Image.open(screenshot) as picture:
                aspect = picture.width / picture.height
            image_height = min(3.75, 8.15 / aspect)
            image_width = image_height * aspect
            picture = slide.shapes.add_picture(str(screenshot), Inches(.65 + (8.15 - image_width) / 2),
                                               Inches(2.42), height=Inches(image_height))
            picture.click_action.hyperlink.address = "https://buffbeefalo.github.io/netmambaplus-reproduction/"
            for n, (label, body) in enumerate(item["cards"]):
                text(slide, label, 9.05, 2.53 + n * 1.76, 3.55, .4, 20, TEAL, True)
                text(slide, body, 9.05, 3.06 + n * 1.76, 3.55, 1.2, 14)
        elif layout == "latency":
            table(slide, ["Batch", "Median batch latency", "95th percentile", "Flows/s from mean"], latency_rows,
                  widths=[1.6, 3.8, 3.2, 3.4], size=18)
            agreement = benchmark["batch_vs_single"]["agreement"]
            text(slide, f"First 128 flows: {percent(agreement)} prediction agreement between single-flow and batched execution.",
                 .8, 5.15, 11.7, .65, 18, TEAL)
        elif layout == "ids":
            values = [("Capture", "Future"), ("Flow cache", "Future"), ("Extractor", "Future"),
                      ("Native tensors", "Measured offline"), ("NetMamba+", "Measured on GPU"), ("Alert policy", "Future")]
            for n, (label, status) in enumerate(values):
                x = .65 + n * 2.04
                done = status != "Future"
                rectangle(slide, x, 3.05, 1.8, 1.5, TEAL if done else "FFFFFF", None if done else MUTED)
                text(slide, label, x + .12, 3.3, 1.56, .68, 15 if label == "NetMamba+" else 17,
                     "FFFFFF" if done else INK, True)
                text(slide, status, x + .12, 4.03, 1.56, .38, 11, "FFFFFF" if done else MUTED)
                if n < 5:
                    text(slide, "→", x + 1.81, 3.55, .25, .4, 18, TEAL)
            text(slide, "Measure independent-traffic accuracy, false alerts, packet loss, flow waiting and queue delay.",
                 .82, 5.15, 11.7, .65, 19)
        else:
            raise ValueError(f"Unknown slide layout: {layout}")
        rectangle(slide, .65, 6.35, 12.03, .56, "E0EAE3")
        text(slide, item["takeaway"], .8, 6.45, 11.74, .4, 12.5, INK)
        footer = text(slide, "NETMAMBA+  /  MEASURED RESEARCH  /  " + source["date"], .65, 7.13, 11.4, .2, 9, MUTED)
        footer.text_frame.paragraphs[0].runs[0].hyperlink.address = REPO + "/tree/main/docs/customer"
        text(slide, f"{index:02d}", 12.2, 7.08, .5, .25, 11, MUTED)
    target = CUSTOMER / "NetMambaPlus-customer-slides.pptx"
    deck.save(target)
    (CUSTOMER / "talk-track.md").write_text("\n".join(notes), encoding="utf-8")
    if shutil.which("soffice"):
        with tempfile.TemporaryDirectory(prefix="netmamba-libreoffice-") as profile:
            subprocess.run(["soffice", f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf",
                            "--outdir", str(CUSTOMER), str(target)], check=True, timeout=120)
    else:
        raise RuntimeError("PowerPoint built, but LibreOffice is required to create and review its PDF")


def inline(value):
    value = html.escape(value)
    value = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r'<a href="\2" color="#007C78">\1</a>', value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", value)
    value = re.sub(r"`([^`]+)`", r'<font name="Mono">\1</font>', value)
    return value


def build_briefing(tokens):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle

    for name, path in [("Body", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                       ("BodyBold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
                       ("Mono", "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")]:
        pdfmetrics.registerFont(TTFont(name, path))
    pdfmetrics.registerFontFamily("Body", normal="Body", bold="BodyBold", italic="Body", boldItalic="BodyBold")
    base = dict(fontName="Body", textColor=colors.HexColor("#" + INK), alignment=TA_LEFT)
    styles = {"body": ParagraphStyle("body", fontSize=10.1, leading=14.3, spaceAfter=8, **base),
              "h1": ParagraphStyle("h1", fontSize=23, leading=28, spaceAfter=17, **base),
              "h2": ParagraphStyle("h2", fontSize=13, leading=18, spaceBefore=11, spaceAfter=7, **base),
              "bullet": ParagraphStyle("bullet", fontSize=10.1, leading=14.1, leftIndent=11, spaceAfter=6, **base),
              "cell": ParagraphStyle("cell", fontSize=8.7, leading=12, **base)}
    resolved = substitute((CUSTOMER / "briefing-source.md").read_text(), tokens)
    (CUSTOMER / "briefing.md").write_text(resolved)
    document = SimpleDocTemplate(str(CUSTOMER / "NetMambaPlus-customer-briefing.pdf"), pagesize=A4,
                                 leftMargin=48, rightMargin=48, topMargin=57, bottomMargin=47,
                                 title="NetMamba+ — what we built and measured", author="NetMamba+ reproduction project")
    story, lines, i = [], resolved.splitlines(), 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "<!-- page -->":
            story.append(PageBreak())
        elif line.startswith("# "):
            story.append(Paragraph(inline(line[2:]), styles["h1"]))
        elif line.startswith("## "):
            story.append(Paragraph(inline(line[3:]), styles["h2"]))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-:]+", cell) for cell in cells):
                    rows.append([Paragraph(inline(cell), styles["cell"]) for cell in cells])
                i += 1
            width = A4[0] - 96
            table = Table(rows, colWidths=[width / len(rows[0])] * len(rows[0]), repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE9E3")),
                                       ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F5F0")]),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                                       ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
            story += [table, Spacer(1, 11)]
            continue
        elif line.startswith("- "):
            story.append(Paragraph("• " + inline(line[2:]), styles["bullet"]))
        else:
            story.append(Paragraph(inline(line), styles["body"]))
        i += 1

    def page(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#" + TEAL))
        canvas.setFont("BodyBold", 8)
        canvas.drawString(48, A4[1] - 30, "NETMAMBA+  /  CUSTOMER RESEARCH BRIEFING")
        canvas.setFillColor(colors.HexColor("#" + MUTED))
        canvas.setFont("Body", 7.5)
        canvas.drawString(48, 25, "Measured source-based experiment · 15 September 2026")
        canvas.drawRightString(A4[0] - 48, 25, str(doc.page))
        canvas.linkURL(REPO + "/tree/main/docs/customer", (48, 21, 405, 35), relative=0)
        canvas.restoreState()
    document.build(story, onFirstPage=page, onLaterPages=page)


def results_markdown(results, benchmark, tokens):
    text = ["# Measured results", "", tokens["results_table"], "",
            "All three runs completed 120 epochs and 7,920 updates each. Checkpoints were selected by validation accuracy. Seed 0 was fixed for the demo before test scores.", "",
            "Paper Table IV reports NetMamba+ CICIoT2022 accuracy/F1 of **97.50%**. The source-based batch/rate settings, GB10 runtime and unresolved pretraining history differ; this is not a controlled equivalence comparison.", "",
            "## Variability across seeds", "", "| Metric | Mean | Sample SD (percentage points) |", "|---|---:|---:|"]
    for key, values in results["aggregate"].items():
        text.append(f"| {key} | {percent(values['mean'])} | {values['sample_standard_deviation'] * 100:.3f} |")
    text += ["", "Three seeds on one split do not establish uncertainty across other networks. All values are independently recomputed from saved predictions; weighted averages use true-class support and macro averages give six classes equal weight.", "",
             "## Per-class seed-0 results", "", "| Class | Support | Precision | Recall | F1 |", "|---|---:|---:|---:|---:|"]
    names = ["Flood", "RTSP Brute Force", "Power–Audio", "Power–Other", "Power–Cameras", "Power–Home Automation"]
    metrics = results["seeds"][0]["metrics"]
    for i, name in enumerate(names):
        text.append(f"| {name} | {metrics['support'][i]} | {percent(metrics['precision_per_class'][i])} | {percent(metrics['recall_per_class'][i])} | {percent(metrics['f1_per_class'][i])} |")
    text += ["", "![Seed-0 confusion matrix](figures/seed0-confusion.png)", "", "Rows are true classes; columns are predictions. Counts sum to 1,041. Prediction errors are retained in the replay.", "",
             "## Training evidence", "", "| Seed | Selected epoch (zero-based) | Best validation accuracy | First → last train loss | Elapsed (shared GPU) |", "|---|---:|---:|---:|---:|"]
    for row in results["seeds"]:
        text.append(f"| {row['seed']} | {row['selected_epoch_zero_based']} | {percent(row['best_validation_accuracy'])} | {row['train_loss_first_epoch']:.4f} → {row['train_loss_last_epoch']:.4f} | {row['training_elapsed_seconds_shared_gpu'] / 60:.2f} min |")
    text += ["", "![Learning curves](figures/learning-curves.png)", "", "The full native epoch logs are retained. Elapsed training times include resource sharing and are not isolated GPU performance measurements. The collector checks the actual selected-checkpoint optimizer counters and changed, finite model parameters.", "",
             "## Strict inference and checkpoint identity", "", "Native final evaluation, separate strict evaluation and independently reconstructed confusion matrices agree for every seed. These re-executions use the same fixed test set and do not constitute new independent holdouts.", "", "| Seed | Selected classifier SHA-256 |", "|---|---|"]
    for row in results["seeds"]:
        text.append(f"| {row['seed']} | `{row['checkpoint_sha256']}` |")
    text += ["", "The model-only exports are tensor-identical and have separate hashes/provenance. Original checkpoints and raw research assets are not silently substituted by the replay.", "",
             "## Model-only latency", "", tokens["latency_table"], "",
             "GB10, float16 autocast, batches 1/16/128, 20 warmups and 100 synchronized measurements per batch. Capture, flow waiting, preprocessing, data loading, transfers and alert handling are excluded. The measurement JSON contains every latency sample and a process snapshot.", "",
             f"Single-flow versus batched prediction agreement on the first 128 test flows: **{percent(benchmark['batch_vs_single']['agreement'])}**. Maximum absolute logit difference: `{benchmark['batch_vs_single']['max_abs_logit_difference']}`.", "",
             "## Limits that accompany the numbers", "",
             "The official split contains five exact stored inputs shared between train/validation and six between train/test. Recorded identifiers and raw fingerprints do not prove capture independence; additional collisions after normalization are possible. Earlier exposure of the released pretraining checkpoint is unknown. These results do not establish unknown-attack performance, calibrated confidence, an operational false-alert rate or NPU/SmartNIC deployment.", "",
             "Primary machine-readable evidence: [results.json](evidence/results.json), [benchmark](evidence/benchmark/metrics.json), [evidence index](evidence/artifact-index.json)."]
    (CUSTOMER / "results.md").write_text("\n".join(text) + "\n")


def main():
    results, benchmark, tokens, rows, latencies = data_and_tokens()
    charts(results)
    results_markdown(results, benchmark, tokens)
    build_briefing(tokens)
    build_slides(results, benchmark, tokens, rows, latencies)
    print("Built briefing PDF, editable PowerPoint, slide PDF, charts, notes and result tables from reviewed evidence.")


if __name__ == "__main__":
    main()
