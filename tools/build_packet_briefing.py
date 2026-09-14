"""Build a PDF, PowerPoint and matching script from completed packet-study evidence."""

import argparse
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import packet_study as study


def content(evidence):
    results = study.read_json(evidence / 'results.json')
    manifest = study.read_json(evidence / 'manifest.json')
    controls = study.read_json(evidence / 'controls/evaluation/results.json')
    if results['status'] != 'completed' or len(results['results']) != 6:
        raise ValueError('Present only a completed six-arm packet study')
    table = [['Training / initialization', 'CIC test', 'UNSW test']]
    for arm in results['results']:
        table.append([arm['arm']['name'].replace('_', ' / '), *[
            f"{arm['tests'][source]['metrics']['group_weighted']['balanced_accuracy']:.2%}"
            for source in ('cic', 'unsw')]])
    rows = manifest['global']['rows']
    metadata_control = next(item for item in controls['results'] if item['name'] == 'unsw_metadata_linear')
    metadata_score = metadata_control['tests']['unsw']['metrics']['group_weighted']['balanced_accuracy']
    return [
        {'title': 'Both CSVs now train NetMamba+', 'subtitle': 'Packet study · addition to the preserved flow reproduction',
         'lines': [f'{rows:,} rows validated across CICIDS2017 and UNSW CSV exports.',
             'Six native GPU training runs: 1,000 supervised updates each.',
             'A joint model learns from both sources; saved checkpoints support unlabeled inference.'],
         'notes': 'Previously these CSVs were only profiled. The original reproduction trained on a separate CICIoT2022 flow release. This addition gives both uploaded CSVs a real role: their stored bytes and supplied binary labels train the original NetMamba+ encoder through a new packet adapter. It does not retroactively change the original experiment or paper result.'},
        {'title': 'A packet is one message fragment', 'subtitle': 'The paper model originally consumed a sequence from a connection',
         'lines': ['Original route: five packet byte segments, twenty sizes and twenty intervals.',
             'CSV route: one row with 1,500 byte slots, four metadata values and a label.',
             'Connection identity and reliable sequence order are missing from these exports.',
             'We retain packet inputs; adjacent rows are never invented into flows.'],
         'notes': 'Think of a flow as a conversation and a packet as one fragment of a message. The paper combines byte content with size and timing sequences. The CSV exports preserve bytes but do not provide the verified joins needed to rebuild those conversations. CICIDS2017 and CICIoT2022 are different datasets. The uploaded PDF explains the architecture; it does not establish that these particular CSV exports were its original training files.'},
        {'title': 'Exactly what goes into the model', 'subtitle': 'Native four-block Mamba encoder, with an explicit packet input contract',
         'lines': ['1,500 integer bytes → float32 values between −1 and +1.',
             'Four bytes per patch → 375 byte tokens + three summary positions.',
             'Size and interval inputs are empty; two prefixes contain no observed flow data.',
             'Native summary pooling → new two-class head → benign / attack logits.'],
         'notes': 'All 1,500 stored slots are retained, including zeros, because the export does not reliably identify padding. TTL, total length, protocol, time delta, source identity and label never enter the native forward call. The native classifier has 378 sequence positions and 1,852,416 parameters. Its original six-class flow head is replaced by a fresh bias-free two-class head. Labels enter only the supervised loss. Softmax scores are uncalibrated, not a probability guarantee of operational danger.'},
        {'title': 'What training actually tests', 'subtitle': 'CIC-only, UNSW-only and joint training; pretrained and scratch for each',
         'lines': ['Batch 64, seed 0, AdamW; same 1,000-update budget and paired batch streams.',
             'Joint batches contain 32 examples from each dataset.',
             'Pretrained: 50 trunk states transferred, 31 decoder states excluded.',
             'Scratch: fresh weights, matching positional coordinates and starting head.',
             'Validation chooses checkpoints; every selection freezes before test inference.'],
         'notes': 'Every model receives 64,000 sampled group presentations, with replacement. A joint model receives 32,000 presentations from each source. This is a fixed-budget experiment, not an epoch over all 1.49 million rows. NetMamba+ learns its position embeddings; pretrained positional weights are part of the transfer treatment. Scratch uses a freshly initialized native 443-position table with the same remap. The comparison has one seed and does not establish a statistically robust improvement.'},
        {'title': 'Measured held-out packet results', 'subtitle': 'Balanced accuracy · each distinct payload has total weight one',
         'table': table, 'lines': [],
         'caveat': f'UNSW-only metadata control: {metadata_score:.2%}. Native NetMamba+ is not universally best here.',
         'notes': f'Balanced accuracy averages benign recall and attack recall. Group weighting prevents repeated identical bytes from dominating the result. Each model is evaluated on both source test partitions. For a single-source model, the opposite-source column is transfer without supervised fitting or validation selection on that source. The joint model has trained on both sources. The UNSW-only metadata control scored {metadata_score:.2%} on UNSW, outperforming the native models there. All scores are specific to these exports, splits, caps and fixed budget; see the report for row weighting, macro-F1, AUROC, average precision, false positives and subtype support.'},
        {'title': 'Why the data checks matter', 'subtitle': 'Repeated bytes and contradictory labels are retained, not hidden',
         'lines': ['478,044 distinct payloads across both full exports.',
             'Ten payloads appear in both files; nine have conflicting benign/attack labels.',
             'Identical inputs stay in the same train, validation or test partition globally.',
             'Contradictory groups contribute their observed label proportions to the loss.',
             'CIC test cap: 20,000 groups; UNSW test: 5,930 groups.'],
         'notes': 'Equal payload bytes do not prove these are the same physical packet or capture. Exact hashing prevents exact-input leakage, but not near duplicates or related captures. Selected CIC test support includes zero PortScan rows and only four DDoS rows, so this study cannot establish coverage for those attacks. Row-weighted results describe represented CSV rows; group-weighted results describe distinct stored inputs. Majority, byte-histogram and metadata-only linear models are diagnostic controls, not replacements for the native-model deliverable.'},
        {'title': 'Use it and check it', 'subtitle': 'Commands and exact arguments are in packet-study-setup.md',
         'lines': ['Prepare both local CSVs with packet_data.py.',
             'Freeze a fresh local protocol, then run train_packet_model.py.',
             'Pass an unlabeled payload CSV and packet checkpoint to predict_packets.py.',
             'Read predictions.jsonl and its fingerprinted completion receipt.',
             'Run review_packet_study.py to recompute the published evidence offline.',
             'All 25,930 predicted classes agree; small GPU score differences remain.'],
         'notes': 'The strict predictor accepts the 1,500 payload columns alone, or those plus the four metadata columns; it rejects a label column. It validates the full file, even when an explicit row cap limits prediction. No-label inference is an actual model call, not replay. All 25,930 predicted classes agreed with the saved joint-model evaluation. Matched FP32 replay still had raw-score differences up to 0.000004411; two CIC rows failed the recorded rtol=1e-4, atol=1e-6 comparison. Those failures remain visible, so no bit-exact numerical claim is made. The public saved records allow CPU arithmetic checks without private raw CSVs. Native GPU execution was measured on GB10; other physical GPUs, NPUs and SmartNICs are not validated by these results.'},
        {'title': 'What you can defend', 'subtitle': 'Present the measured addition with its limits',
         'lines': ['Both supplied CSVs genuinely train the native NetMamba+ packet adaptation.',
             'The PDF-guided original flow reproduction remains a separate experiment.',
             'Missing: live capture pipeline, verified flow joins and customer-network validation.',
             'Missing: calibrated packet confidence, NPU export and SmartNIC deployment.',
             'Existing v4 video covers flows and calibration; send this addendum with it.'],
         'notes': 'A simple IDS prototype could feed already-formatted unlabeled packet rows into the saved joint checkpoint and log predictions for analyst review. It currently does not capture, block or inspect live traffic. An eventual SmartNIC pipeline would need packet extraction, batching and a supported numerical implementation of the Mamba scan, followed by end-to-end accuracy, throughput and latency tests. The council was consulted but ended ESCALATED and UNRATIFIED; primary-source and runtime checks resolved implementation issues directly. Do not call that consensus or claim the paper’s 97.50 percent result was reproduced.'},
    ]


def script_text(slides):
    script = '# NetMamba+ packet addendum — presenter script\n\nThis follows the eight-slide PDF and PowerPoint. It is a later addition to the preserved v4 course.\n\n'
    for index,item in enumerate(slides,1):
        script += f"## Slide {index}: {item['title']}\n\n{item['notes']}\n\n"
    return script.rstrip()+'\n'


def build(evidence, output):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import textwrap
    slides = content(evidence)
    output.mkdir(parents=True, exist_ok=True)
    font = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    pdfmetrics.registerFont(TTFont('PacketSans', str(font)))
    deck = Presentation(); deck.slide_width = Inches(13.333333); deck.slide_height = Inches(7.5)
    pdf_path = output / 'NetMambaPlus-packet-addendum.pdf'
    ppt_path = output / 'NetMambaPlus-packet-addendum.pptx'
    pdf = canvas.Canvas(str(pdf_path), pagesize=(960,540), invariant=1)
    pdf.setTitle('NetMamba+ packet study — measured CSV integration')
    navy, blue, gray = '101D35', '63A9FF', 'B9C8DE'
    def ppt_text(slide, text, x, y, width, height, size, color='FFFFFF', bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame; frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        for i, line in enumerate(text.split('\n')):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line; paragraph.font.name = 'DejaVu Sans'; paragraph.font.size = Pt(size)
            paragraph.font.bold = bold; paragraph.font.color.rgb = RGBColor.from_string(color)
    for index, item in enumerate(slides,1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = RGBColor.from_string(navy)
        ppt_text(slide, 'NETMAMBA+  /  PACKET STUDY', .6,.35,12,.3,11,blue,True)
        ppt_text(slide, item['title'], .6,.95,12, .7,29,bold=True)
        ppt_text(slide, item['subtitle'], .6,1.75,12,.6,15,gray)
        pdf.setFillColor('#'+navy); pdf.rect(0,0,960,540,fill=1,stroke=0)
        pdf.setFont('PacketSans',11);pdf.setFillColor('#'+blue);pdf.drawString(43,500,'NETMAMBA+  /  PACKET STUDY')
        pdf.setFont('PacketSans',29);pdf.setFillColor('#FFFFFF');pdf.drawString(43,438,item['title'])
        pdf.setFont('PacketSans',15);pdf.setFillColor('#'+gray);pdf.drawString(43,400,item['subtitle'])
        if 'table' in item:
            y = 2.55
            for row_index, row in enumerate(item['table']):
                for cell, x, width in zip(row,(.6,7.5,10),(6.5,2,2)):
                    ppt_text(slide,cell,x,y,width,.45,16 if row_index else 14,blue if row_index == 0 else 'FFFFFF',row_index == 0)
                    pdf.setFont('PacketSans',14 if row_index == 0 else 16)
                    pdf.setFillColor('#'+(blue if row_index == 0 else 'FFFFFF'))
                    pdf.drawString(x*72,540-y*72-20,cell)
                y += .46
            ppt_text(slide,item['caveat'],.6,6.18,12,.5,14,'FFD384')
            pdf.setFont('PacketSans',14);pdf.setFillColor('#FFD384');pdf.drawString(43,80,item['caveat'])
        else:
            y = 2.65
            for line in item['lines']:
                wrapped = textwrap.wrap(line, width=83)
                ppt_text(slide, '\n'.join(wrapped),.65,y,12,.8,19)
                pdf.setFont('PacketSans',19);pdf.setFillColor('#FFFFFF')
                for j, part in enumerate(wrapped):pdf.drawString(47,540-y*72-21-j*25,part)
                y += .5 + .34*(len(wrapped)-1)
        footer=f'Packet adaptation · GB10 evidence · original flow results unchanged                                      {index} / {len(slides)}'
        ppt_text(slide,footer,.6,7.02,12,.3,10,gray)
        pdf.setFont('PacketSans',10);pdf.setFillColor('#'+gray);pdf.drawString(43,24,footer)
        slide.notes_slide.notes_text_frame.text=item['notes']
        pdf.showPage()
    deck.save(ppt_path);pdf.save()
    (output/'packet-addendum-script.md').write_text(script_text(slides),encoding='utf-8')
    (output/'packet-addendum-source.json').write_text(json.dumps(slides,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    print(json.dumps({'slides':len(slides),'pdf':str(pdf_path),'pptx':str(ppt_path)}))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--evidence',type=Path,default=ROOT/'docs/customer/evidence/packet-study')
    parser.add_argument('--output',type=Path,default=ROOT/'docs/customer/packet-addendum')
    args=parser.parse_args();build(args.evidence,args.output)


if __name__=='__main__':main()
