"""Packet publication uses its recorded scope without rewriting archived copy."""

from contextlib import redirect_stderr, redirect_stdout
import hashlib
import html
from importlib.util import find_spec
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'tools'))
import build_video_documents as documents
import build_video_page as page
import tests.test_course_video as fixtures


PACKET_TITLE = 'Packet CSVs & measured evidence'
PACKET_SUMMARY = 'Both CSVs train one packet model; <scores> retain their measured limits.'
DOCUMENT_TOOLS = (find_spec('pptx') is not None and find_spec('reportlab') is not None
                  and find_spec('PIL') is not None and shutil.which('pdftotext')
                  and shutil.which('pdfinfo')
                  and Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf').is_file())


def packet_manifest():
    manifest = fixtures.VideoTests().hour_manifest()
    manifest.update(title=PACKET_TITLE, summary=PACKET_SUMMARY,
                    release_tag='course-video-v5', scheduled_seconds=2700)
    for index, chapter in enumerate(manifest['chapters']):
        chapter.update(start=index * 900, end=(index + 1) * 900)
        speech, pause = manifest['scenes'][index * 2:index * 2 + 2]
        speech.update(start=chapter['start'], end=chapter['end'] - 20)
        pause.update(start=chapter['end'] - 20, end=chapter['end'])
    return manifest


def publication_fixture(root, version='v5', *, summary=True):
    manifest = packet_manifest()
    manifest['release_tag'] = 'course-video-' + version
    if not summary:
        manifest.pop('summary')
    source = fixtures.VideoTests().hour_source()
    source.update(title=manifest['title'], release_tag=manifest['release_tag'],
                  target_seconds=2700, extra_references={'notes': 'README.md'},
                  fact_bindings_sha256=hashlib.sha256(b'{}').hexdigest())
    if summary:
        source['summary'] = manifest['summary']
    for chapter in source['chapters']:
        chapter['seconds'] = 900
    source_path = root / f'docs/customer/video-course-{version}-source.json'
    source_path.parent.mkdir(parents=True, exist_ok=True)
    source_path.write_text(json.dumps(source), encoding='utf-8')
    (source_path.parent / 'course-source.json').write_text(
        json.dumps({'facts': {}, 'references': {}}), encoding='utf-8')
    (root / 'README.md').write_text('Fixture reference\n', encoding='utf-8')
    (root / 'docs/repository-walkthrough.md').write_text(
        '# File guide\n\nThe fixture appendix describes the supplied files.\n',
        encoding='utf-8')
    manifest.update(source_path=source_path.relative_to(root).as_posix(),
                    source_sha256=page.digest(source_path))
    manifest['references']['notes']['sha256'] = page.digest(root / 'README.md')
    media = root / 'media' / version
    media.mkdir(parents=True)
    (media / 'media-manifest.json').write_text(json.dumps(manifest), encoding='utf-8')
    return source_path, media, manifest


class PacketPublicationTests(unittest.TestCase):
    def test_packet_page_uses_escaped_summary_title_and_measured_duration(self):
        rendered = page.render(packet_manifest())
        self.assertIn('<title>' + html.escape(PACKET_TITLE) + '</title>', rendered)
        self.assertIn('<meta name="description" content="' + html.escape(PACKET_SUMMARY) + '">', rendered)
        self.assertIn('<p class="status">' + html.escape(PACKET_SUMMARY) + '</p>', rendered)
        self.assertIn('One continuous 45:00 video', rendered)
        self.assertIn('src="v5/measured-course.mp4"', rendered)
        self.assertIn('video-verification-v5.md', rendered)
        for obsolete in ('86.65%', '97.50%', 'Later addition:', 'original flow experiment'):
            self.assertNotIn(obsolete, rendered)
        self.assertNotIn('<scores>', rendered)

    def test_packet_page_requires_its_summary_instead_of_using_flow_copy(self):
        manifest = packet_manifest()
        manifest.pop('summary')
        with self.assertRaisesRegex(ValueError, 'summary'):
            page.render(manifest)

    def test_packet_documents_require_summary_before_loading_renderers(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, media, _ = publication_fixture(root, summary=False)
            with (patch.object(documents, 'ROOT', root),
                  self.assertRaisesRegex(ValueError, 'summary')):
                documents.build(root / 'work', media, source)

    def test_default_render_reads_the_packet_media_manifest(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            packet = packet_manifest()
            archived = fixtures.VideoTests().hour_manifest()
            archived.update(title='Archived flow course', release_tag='course-video-v4')
            for version, manifest in [('v4', archived), ('v5', packet)]:
                media = root / version
                media.mkdir()
                (media / 'media-manifest.json').write_text(json.dumps(manifest), encoding='utf-8')
            selected = root / page.DESTINATION.name
            with patch.object(page, 'DESTINATION', selected):
                rendered = page.render()
            self.assertIn('<h1>' + html.escape(PACKET_TITLE) + '</h1>', rendered)
            self.assertNotIn('Archived flow course', rendered)

    def test_archived_pages_retain_their_original_scope(self):
        for version in ('v3', 'v4'):
            with self.subTest(version=version):
                manifest = fixtures.VideoTests().hour_manifest()
                manifest['release_tag'] = 'course-video-' + version
                rendered = page.render(manifest)
                self.assertIn('Actual three-run mean: <strong>86.65%</strong>', rendered)
                self.assertIn('97.50%', rendered)
                self.assertIn(f'src="{version}/measured-course.mp4"', rendered)
                self.assertEqual('Later addition:' in rendered, version == 'v4')

    def test_cli_can_publish_v5_to_the_active_page(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, media, _ = publication_fixture(root)
            active = root / 'index.html'
            arguments = ['build_video_page.py', '--source', str(source), '--media-dir', str(media)]
            with (patch.object(page, 'ROOT', root), patch.object(page, 'WATCH_PAGE', active),
                  patch.object(sys, 'argv', arguments), redirect_stdout(io.StringIO())):
                page.main()
            self.assertIn(html.escape(PACKET_SUMMARY), active.read_text(encoding='utf-8'))

    def test_cli_cannot_replace_active_packet_page_with_archived_v4(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, media, _ = publication_fixture(root, 'v4', summary=False)
            active = root / 'index.html'
            active.write_text('Preserved packet page', encoding='utf-8')
            arguments = ['build_video_page.py', '--source', str(source), '--media-dir', str(media)]
            with (patch.object(page, 'ROOT', root), patch.object(page, 'WATCH_PAGE', active),
                  patch.object(sys, 'argv', arguments), redirect_stderr(io.StringIO()),
                  redirect_stdout(io.StringIO()), self.assertRaises(SystemExit)):
                page.main()
            self.assertEqual(active.read_text(encoding='utf-8'), 'Preserved packet page')
            archived = root / 'archive.html'
            with (patch.object(page, 'ROOT', root), patch.object(page, 'WATCH_PAGE', active),
                  patch.object(sys, 'argv', arguments + ['--output', str(archived)]),
                  redirect_stdout(io.StringIO())):
                page.main()
            self.assertIn('86.65%', archived.read_text(encoding='utf-8'))
            self.assertEqual(active.read_text(encoding='utf-8'), 'Preserved packet page')

    @unittest.skipUnless(DOCUMENT_TOOLS, 'Optional document renderer and PDF tools are unavailable')
    def test_packet_companions_use_selected_title_summary_and_actual_duration(self):
        from PIL import Image
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, media, manifest = publication_fixture(root)
            work = root / 'work'
            (work / 'frames').mkdir(parents=True)
            (media / 'frames').mkdir()
            for scene in manifest['scenes']:
                frame = work / 'frames' / (scene['id'] + '.png')
                Image.new('RGB', (96, 54), '#e2eeeb').save(frame)
                published = media / 'frames' / frame.name
                published.write_bytes(frame.read_bytes())
                scene['visuals'] = [{'path': published.relative_to(root).as_posix(),
                                     'sha256': page.digest(frame)}]
            (media / 'media-manifest.json').write_text(json.dumps(manifest), encoding='utf-8')
            with patch.object(documents, 'ROOT', root), redirect_stdout(io.StringIO()):
                documents.build(work, media, source)
            deck = Presentation(media / documents.SLIDES)
            self.assertEqual(deck.core_properties.title, PACKET_TITLE)
            self.assertEqual(len(deck.slides), 6)
            text = subprocess.check_output(['pdftotext', str(media / documents.HANDBOOK), '-'], text=True)
            self.assertIn(PACKET_TITLE, text)
            self.assertIn(PACKET_SUMMARY, ' '.join(text.split()))
            self.assertIn('45:00', text)
            links = subprocess.check_output(
                ['pdfinfo', '-url', str(media / documents.HANDBOOK)], text=True)
            link_rows = [line.split() for line in links.splitlines()]
            for repository in ('netmambaplus-reproduction', 'netmambaplus-client'):
                self.assertIn(['1', 'Annotation',
                               'https://github.com/buffbeefalo/' + repository], link_rows)
            for obsolete in ('hour-long', '86.65%', '97.50%'):
                self.assertNotIn(obsolete, text)
            for name in (documents.HANDBOOK, documents.SLIDE_PDF):
                info = subprocess.check_output(['pdfinfo', str(media / name)], text=True)
                self.assertIn(PACKET_TITLE, info)
            published_manifest = json.loads((media / 'media-manifest.json').read_text())
            for name in (documents.SLIDES, documents.SLIDE_PDF, documents.HANDBOOK,
                         'companion-manifest.json'):
                self.assertEqual(published_manifest['artifacts'][name]['sha256'],
                                 page.digest(media / name))
            rendered = page.render(published_manifest, media_dir=media, output=root / 'index.html')
            for name in (documents.SLIDES, documents.SLIDE_PDF, documents.HANDBOOK):
                self.assertIn('href="media/v5/' + name + '"', rendered)

    @unittest.skipUnless(DOCUMENT_TOOLS, 'Optional document renderer and PDF tools are unavailable')
    def test_archived_companions_keep_the_preserved_title_and_flow_summary(self):
        from PIL import Image
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, media, manifest = publication_fixture(root, 'v4', summary=False)
            work = root / 'work'
            (work / 'frames').mkdir(parents=True)
            (media / 'frames').mkdir()
            for scene in manifest['scenes']:
                frame = work / 'frames' / (scene['id'] + '.png')
                Image.new('RGB', (96, 54), '#e2eeeb').save(frame)
                published = media / 'frames' / frame.name
                published.write_bytes(frame.read_bytes())
                scene['visuals'] = [{'path': published.relative_to(root).as_posix(),
                                     'sha256': page.digest(frame)}]
            (media / 'media-manifest.json').write_text(json.dumps(manifest), encoding='utf-8')
            with patch.object(documents, 'ROOT', root), redirect_stdout(io.StringIO()):
                documents.build(work, media, source)
            self.assertEqual(Presentation(media / documents.SLIDES).core_properties.title,
                             'NetMamba+ — complete repository course')
            text = subprocess.check_output(['pdftotext', str(media / documents.HANDBOOK), '-'], text=True)
            self.assertIn('single hour-long video', text)
            self.assertIn('86.65%', text)
            self.assertIn('97.50%', text)
            self.assertNotIn('Project repositories:', text)
            self.assertEqual(json.loads((media / 'media-manifest.json').read_text()), manifest)


if __name__ == '__main__':
    unittest.main()
